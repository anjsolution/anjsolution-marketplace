"""장애 history 워크북(.xlsm/.xlsx) 입출력 — 열 정의, 행 읽기, 셀 한 칸 수정.

다른 history_* 스크립트가 공용으로 쓰는 라이브러리다. 직접 실행하지 않는다.
셀 수정은 시트 XML의 해당 셀만 바꾼다. openpyxl로 통째로 다시 저장하면 통계 시트의
차트·매크로가 손상되므로 쓰지 않는다.
"""
import os
import re
import shutil
import sys
import zipfile
from datetime import datetime
from xml.sax.saxutils import escape

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from history_rules import kst, text  # noqa: E402,F401 — 다른 스크립트가 여기서 함께 가져간다

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl이 필요합니다: python -m pip install openpyxl")

SHEET = "장애 history"

# 열 문자 → (내부 키, 원본 추출 JSON의 키). N열(복구 소요시간)은 계산 열이라 읽지 않는다.
COLUMNS = {
    "A": ("no", "A_접수번호"), "B": ("done", "B_처리완료"), "D": ("code", "D_장애코드"),
    "E": ("received", "E_접수일자"), "F": ("division", "F_본부"), "G": ("tunnel", "G_터널명"),
    "H": ("receiver", "H_접수자"), "I": ("reporter_org", "I_신고자소속"),
    "J": ("reporter", "J_신고자"), "K": ("action", "K_조치코드"), "L": ("resolved", "L_복구일자"),
    "M": ("handler", "M_처리자"), "O": ("content", "O_장애내용"), "P": ("progress", "P_대응경과"),
    "Q": ("cause", "Q_원인"), "R": ("measure", "R_조치"), "S": ("remarks", "S_비고"),
}
SOURCE_KEY = {key: label for key, label in COLUMNS.values()}


def col_index(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def load_rows(path):
    """접수번호(A열)가 숫자인 행만 {내부 키: 셀 값, excel_row} 목록으로 읽는다."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    last = max(col_index(c) for c in COLUMNS)
    rows = []
    for excel_row, values in enumerate(ws.iter_rows(min_row=2, max_col=last, values_only=True), start=2):
        row = {key: values[col_index(c) - 1] for c, (key, _) in COLUMNS.items()}
        if isinstance(row["no"], (int, float)) and not isinstance(row["no"], bool):
            row["no"], row["excel_row"] = int(row["no"]), excel_row
            rows.append(row)
    wb.close()
    return rows


def parse_numbers(specs):
    """'1940-1942', '5,7' 같은 인자를 접수번호 목록으로."""
    numbers = []
    for spec in specs:
        for part in str(spec).split(","):
            if "-" in part:
                lo, hi = (int(x) for x in part.split("-", 1))
                numbers.extend(range(lo, hi + 1))
            elif part.strip():
                numbers.append(int(part))
    return numbers


def write_cells(path, cells):
    """{"K1944": "ORCH"} 형태로 받은 셀만 문자열로 바꾸고 백업 경로를 돌려준다."""
    path = str(path)
    lock = os.path.join(os.path.dirname(os.path.abspath(path)), "~$" + os.path.basename(path))
    if os.path.exists(lock):
        sys.exit("엑셀에서 파일이 열려 있습니다. 닫은 뒤 다시 실행하세요.")
    with zipfile.ZipFile(path) as z:
        infos, data = z.infolist(), {i.filename: z.read(i.filename) for i in z.infolist()}
    wb = data["xl/workbook.xml"].decode("utf-8")
    sheet = re.search(r'<sheet\b[^>]*\bname="%s"[^>]*>' % re.escape(escape(SHEET)), wb)
    rid = re.search(r'r:id="([^"]+)"', sheet.group(0)).group(1) if sheet else "rId1"
    rels = data["xl/_rels/workbook.xml.rels"].decode("utf-8")
    target = re.search(r'<Relationship\b[^>]*Id="%s"[^>]*>' % rid, rels).group(0)
    part = "xl/" + re.search(r'Target="/?(?:xl/)?([^"]+)"', target).group(1)
    xml = data[part].decode("utf-8")
    for ref, value in cells.items():
        xml = _set_cell(xml, ref, value)
    data[part] = xml.encode("utf-8")
    if "fullCalcOnLoad" not in wb:  # 통계 시트가 열 때 다시 계산되도록
        data["xl/workbook.xml"] = re.sub(
            r"<calcPr\b", '<calcPr fullCalcOnLoad="1"', wb, count=1).encode("utf-8")
    backup = f"{path}.bak-{datetime.now():%y%m%d_%H%M%S}"
    shutil.copy2(path, backup)
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w") as out:
        for info in infos:
            out.writestr(info, data[info.filename])
    os.replace(tmp, path)
    return backup


def _set_cell(xml, ref, value):
    row_no = re.search(r"\d+", ref).group(0)
    row = re.search(r'<row\b[^>]*\br="%s"[^>]*>.*?</row>' % row_no, xml, re.S)
    if not row:
        raise SystemExit(f"{ref} 행을 시트에서 찾지 못했습니다.")
    body = row.group(0)
    old = re.search(r'<c\b[^>]*\br="%s"[^>]*?(?:/>|>.*?</c>)' % ref, body, re.S)
    style = re.search(r'\bs="\d+"', old.group(0)) if old else None
    new = (f'<c r="{ref}"{" " + style.group(0) if style else ""} t="inlineStr">'
           f"<is><t>{escape(value)}</t></is></c>")
    if old:
        body = body[:old.start()] + new + body[old.end():]
    else:
        letters = re.match(r"[A-Z]+", ref).group(0)
        later = [m for m in re.finditer(r'<c\b[^>]*\br="([A-Z]+)\d+"', body)
                 if col_index(m.group(1)) > col_index(letters)]
        at = later[0].start() if later else body.rindex("</row>")
        body = body[:at] + new + body[at:]
    return xml[:row.start()] + body + xml[row.end():]
