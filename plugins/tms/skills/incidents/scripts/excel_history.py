"""장애 history 엑셀(.xlsm/.xlsx)을 원격 TMS와 비교해 차이·특이사항을 찾고 등록용 JSON 초안을 만든다.

  python excel_history.py pending <엑셀> --remote 최근조회.json [--after 1944]
  python excel_history.py extract <엑셀> 1945-1950
  python excel_history.py fix     <엑셀> 1945 --action ORCH [--code HL100]

pending : 원격 최근 조회(search_incidents 응답 저장본)에서 마지막 접수번호를 찾아
          그 이후 엑셀에만 있는 행, 접수번호 없는 원격 장애(중복 확인), 조회 범위 안에서
          엑셀 쪽 대응·조치가 더 많은 건을 보고한다.
extract : 접수번호별 장애·대응·조치 초안. review가 비어 있어야 등록할 수 있다.
fix     : 장애코드(D열)·조치코드(K열)를 검증 후 그 셀만 고친다. 백업을 먼저 만든다.

대상 ID는 만들지 않는다. 본부·터널명 원문(target_text)을 catalog로 해석한다.
코드가 도구 값에 맞지 않으면 추측하지 않고 review에 남긴다.
"""
import argparse
import json
import os
import re
import shutil
import sys
import zipfile
from datetime import datetime
from xml.sax.saxutils import escape

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl이 필요합니다: python -m pip install openpyxl")

SHEET = "장애 history"
COLS = {1: "no", 2: "done", 4: "code", 5: "received", 6: "division", 7: "tunnel",
        8: "receiver", 9: "reporter_org", 10: "reporter", 11: "action", 12: "resolved",
        13: "handler", 15: "content", 16: "progress", 17: "cause", 18: "measure", 19: "remarks"}

CATEGORY = {"S": "SOFTWARE", "H": "HARDWARE", "N": "NETWORK", "E": "ENVIRONMENTAL"}
LOCATION = {"L": "ADMIN_SERVER", "M": "TUNNEL_ADMIN", "R": "TUNNEL_DEVICE"}
DEVICE = {"1": "SERVER", "2": "NETWORK", "3": "CONTROLLER", "4": "OTHER"}
SUBTYPE = {"1": "VMS", "2": "LCS", "3": "VDS", "4": "CCTV", "5": "BROADCAST",
           "6": "EMERGENCY_PHONE", "7": "MANUAL_FIRE", "8": "OTHER"}
ACTION = (("target", {"S": "SERVER", "N": "NETWORK", "M": "CONTROLLER", "O": "OTHER"}),
          # O(Other, 운영미숙 포함)는 서버의 조치 방법 OTHER 추가를 전제로 허용한다
          ("method", {"E": "CONFIG", "R": "REBOOT", "P": "PATCH", "C": "REPLACE", "F": "REPAIR",
                      "O": "OTHER"}),
          ("access", {"A": "WIRED", "B": "REMOTE", "C": "ONSITE"}),
          ("type", {"H": "HARDWARE", "S": "SOFTWARE"}))

# "15:59 내용", "6/21 16:48 내용", "24/12/31 17:24 내용", "- 14:46 내용"
LINE_TIME = re.compile(
    r"^[\s\-○•·*]*(?:(?:(?P<y>\d{4}|\d{2})[-./])?(?P<m>\d{1,2})[-./](?P<d>\d{1,2})\s+)?"
    r"(?P<H>\d{1,2}):(?P<M>\d{2})\s*(?P<text>.*)$")
EMPTY_PROGRESS = {"없음", "-", "ㅇㅇ"}


def text(value):
    value = None if value is None else str(value).strip()
    return value or None


def kst(dt):
    return dt.strftime("%Y-%m-%dT%H:%M:%S+09:00") if isinstance(dt, datetime) else None


def load_rows(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    rows = []
    for excel_row, values in enumerate(
            ws.iter_rows(min_row=2, max_col=max(COLS), values_only=True), start=2):
        row = {key: values[i - 1] for i, key in COLS.items()}
        if isinstance(row["no"], (int, float)) and not isinstance(row["no"], bool):
            row["no"], row["excel_row"] = int(row["no"]), excel_row
            rows.append(row)
    wb.close()
    return rows


def parse_incident_code(raw):
    """장애코드 → (값, 오류). 여러 코드면 첫 번째를 쓴다."""
    code = (text(raw) or "").upper().split(",")[0].strip()
    out = {"category": CATEGORY.get(code[:1]), "location": LOCATION.get(code[1:2]),
           "deviceCategory": DEVICE.get(code[2:3]), "controllerSubType": None}
    if out["deviceCategory"] == "CONTROLLER":
        out["controllerSubType"] = SUBTYPE.get(code[3:4])
    bad = not re.fullmatch(r"[A-Z]{2}\d{3}", code) or None in (
        out["category"], out["location"], out["deviceCategory"]) or (
        out["deviceCategory"] == "CONTROLLER" and not out["controllerSubType"])
    return out, (f"D열 장애코드 '{text(raw)}' 해석 불가" if bad else None)


def parse_action_code(raw):
    """조치코드 4글자 → (값, 오류)."""
    code = (text(raw) or "").upper()
    out = {key: table.get(code[i:i + 1]) if len(code) > i else None
           for i, (key, table) in enumerate(ACTION)}
    missing = [key for key, value in out.items() if value is None]
    if len(code) != 4 or missing:
        return out, f"K열 조치코드 '{text(raw)}'의 {', '.join(missing) or '길이'} 해석 불가"
    return out, None


def split_progress(raw, received, handler, review):
    """대응 경과를 줄 단위로 나눈다. 시각 없는 줄은 앞 기록에 이어 붙인다."""
    responses = []
    if text(raw) in EMPTY_PROGRESS:
        return responses
    for line in (text(raw) or "").splitlines():
        if not line.strip():
            continue
        m = LINE_TIME.match(line)
        if m and received:
            year = int(m["y"]) if m["y"] else received.year
            year += 2000 if year < 100 else 0
            month = int(m["m"]) if m["m"] else received.month
            day = int(m["d"]) if m["d"] else received.day
            if not m["y"] and month == 12 and received.month == 1:
                year -= 1  # 연초 접수건의 전년 12월 경과
            try:
                at = kst(datetime(year, month, day, int(m["H"]), int(m["M"])))
            except ValueError:
                review.append(f"P열 대응 경과 시각 오류: '{line.strip()}'")
                at = None
            responses.append({"responseAt": at, "content": text(m["text"]) or line.strip(),
                              "responderName": handler})
        elif responses:
            responses[-1]["content"] += "\n" + line.strip()
        else:
            responses.append({"responseAt": None, "content": line.strip(), "responderName": handler})
            review.append("P열 대응 경과 첫 줄에 시각 없음")
    return responses


def build(row):
    review = []
    handler = text(row["handler"])
    received = row["received"] if isinstance(row["received"], datetime) else None
    if not received:
        review.append("E열 접수일자 없음")
    codes, err = parse_incident_code(row["code"])
    review += [err] if err else []
    incident = {"legacyReceiptNo": row["no"], "receivedAt": kst(received),
                "receiverName": text(row["receiver"]), "reporterName": text(row["reporter"]),
                "reporterOrganization": text(row["reporter_org"]), **codes,
                "memo": text(row["content"])}
    if not incident["memo"]:
        review.append("O열 장애 내용 없음")
    resolution = None
    if text(row["action"]) or text(row["measure"]) or isinstance(row["resolved"], datetime):
        action, err = parse_action_code(row["action"])
        review += [err] if err else []
        resolution = {**action, "resolvedAt": kst(row["resolved"]), "resolverName": handler,
                      "cause": text(row["cause"]), "memo": text(row["measure"]),
                      "remarks": text(row["remarks"])}
    responses = split_progress(row["progress"], received, handler, review)
    return {"receiptNo": row["no"], "ready": not review,
            "target_text": {"본부": text(row["division"]), "터널명": text(row["tunnel"])},
            "incident": incident, "responses": responses,
            "resolution": resolution, "review": review}


def load_remote(path):
    """search_incidents 응답(한 페이지·페이지 배열·목록 배열)에서 장애 목록만 모은다."""
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    items = []
    for page in data if isinstance(data, list) else [data]:
        items.extend(page.get("목록", []) if "목록" in page else [page])
    return items


def pending(rows, remote_path, after):
    items = load_remote(remote_path)
    if after is None and not any("접수번호" in it for it in items):
        sys.exit("원격 조회 결과에 접수번호 필드가 없습니다. MCP 배포 전이면 --after로 마지막 접수번호를 지정하세요.")
    numbered = [it for it in items if it.get("접수번호") is not None]
    last = max(numbered, key=lambda it: it["접수번호"], default=None)
    if after is None:
        if not last:
            sys.exit("조회한 원격 장애에 접수번호가 하나도 없습니다. 조회 범위를 넓히세요.")
        after = last["접수번호"]
    todo = [build(r) for r in rows if r["no"] > after]
    numbers = [d["receiptNo"] for d in todo]
    dup_check = []
    for it in items:
        if it.get("접수번호") is not None:
            continue
        day = (it.get("접수일시") or "")[:10]
        same = [d["receiptNo"] for d in todo if (d["incident"]["receivedAt"] or "")[:10] == day]
        exact = [d["receiptNo"] for d in todo
                 if (d["incident"]["receivedAt"] or "")[:16] == (it.get("접수일시") or "")[:16]]
        if same:
            dup_check.append({"원격_장애코드": it.get("장애코드"), "접수일시": it.get("접수일시"),
                              "메모": (it.get("메모") or "")[:80], "같은날_엑셀": same,
                              "같은시각_엑셀": exact})
    by_no = {r["no"]: r for r in rows}
    additions = []  # 원격 조회 범위 안에서 엑셀 쪽 대응·조치가 더 많은 건
    for it in numbered:
        row = by_no.get(it["접수번호"])
        if not row:
            continue
        draft = build(row)
        remote_count = it.get("대응기록수") or 0
        if len(draft["responses"]) > remote_count:
            additions.append({"접수번호": it["접수번호"], "장애코드": it.get("장애코드"), "추가": "대응",
                              "엑셀": len(draft["responses"]), "원격": remote_count})
        if draft["resolution"] and it.get("조치완료") is False:
            additions.append({"접수번호": it["접수번호"], "장애코드": it.get("장애코드"), "추가": "조치"})
    return {
        "기준_접수번호": after,
        "원격_마지막_장애": last and {k: last.get(k) for k in ("접수번호", "장애코드", "접수일시", "메모")},
        "엑셀_마지막_접수번호": max((r["no"] for r in rows), default=None),
        "결번": sorted(set(range(after + 1, max(numbers) + 1)) - set(numbers)) if numbers else [],
        "등록_가능": [summary(d) for d in todo if d["ready"]],
        "수정_필요": [summary(d) | {"사유": d["review"]} for d in todo if not d["ready"]],
        "중복_확인": dup_check,
        "대응조치_추가분": additions,
    }


def summary(d):
    return {"접수번호": d["receiptNo"], "접수일시": d["incident"]["receivedAt"],
            "터널명": d["target_text"]["터널명"], "내용": (d["incident"]["memo"] or "")[:60]}


def parse_numbers(specs):
    numbers = []
    for spec in specs:
        for part in spec.split(","):
            if "-" in part:
                lo, hi = (int(x) for x in part.split("-", 1))
                numbers.extend(range(lo, hi + 1))
            elif part.strip():
                numbers.append(int(part))
    return numbers


def write_cells(path, cells):
    """시트 XML에서 지정 셀만 문자열로 바꾼다. 차트·매크로 등 나머지 파일은 그대로 복사한다."""
    lock = os.path.join(os.path.dirname(os.path.abspath(path)), "~$" + os.path.basename(path))
    if os.path.exists(lock):
        sys.exit("엑셀에서 파일이 열려 있습니다. 닫은 뒤 다시 실행하세요.")
    with zipfile.ZipFile(path) as z:
        infos, data = z.infolist(), {i.filename: z.read(i.filename) for i in z.infolist()}
    wb = data["xl/workbook.xml"].decode("utf-8")
    sheet = re.search(r'<sheet\b[^>]*\bname="%s"[^>]*>' % re.escape(escape(SHEET)), wb)
    rid = re.search(r'r:id="([^"]+)"', sheet.group(0)).group(1) if sheet else "rId1"
    rels = data["xl/_rels/workbook.xml.rels"].decode("utf-8")
    target = re.search(r'<Relationship\b[^>]*Id="%s"[^>]*>' % rid, rels).group(0)
    part = "xl/" + re.search(r'Target="/?(?:xl/)?([^"]+)"', target).group(1)
    xml = data[part].decode("utf-8")
    for ref, value in cells.items():
        xml = _set_cell(xml, ref, value)
    data[part] = xml.encode("utf-8")
    if "fullCalcOnLoad" not in wb:  # 통계 시트가 열 때 다시 계산되도록
        data["xl/workbook.xml"] = re.sub(r"<calcPr\b", '<calcPr fullCalcOnLoad="1"', wb, count=1).encode("utf-8")
    backup = f"{path}.bak-{datetime.now():%y%m%d_%H%M%S}"
    shutil.copy2(path, backup)
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w") as out:
        for info in infos:
            out.writestr(info, data[info.filename])
    os.replace(tmp, path)
    return backup


def _col_index(ref):
    n = 0
    for ch in re.match(r"[A-Z]+", ref).group(0):
        n = n * 26 + ord(ch) - 64
    return n


def _set_cell(xml, ref, value):
    row_no = re.search(r"\d+", ref).group(0)
    row = re.search(r'<row\b[^>]*\br="%s"[^>]*>.*?</row>' % row_no, xml, re.S)
    if not row:
        raise SystemExit(f"{ref} 행을 시트에서 찾지 못했습니다.")
    body = row.group(0)
    old = re.search(r'<c\b[^>]*\br="%s"[^>]*?(?:/>|>.*?</c>)' % ref, body, re.S)
    style = re.search(r'\bs="\d+"', old.group(0)) if old else None
    new = f'<c r="{ref}"{" " + style.group(0) if style else ""} t="inlineStr"><is><t>{escape(value)}</t></is></c>'
    if old:
        body = body[:old.start()] + new + body[old.end():]
    else:
        later = [m for m in re.finditer(r'<c\b[^>]*\br="([A-Z]+)\d+"', body)
                 if _col_index(m.group(1)) > _col_index(ref)]
        at = later[0].start() if later else body.rindex("</row>")
        body = body[:at] + new + body[at:]
    return xml[:row.start()] + body + xml[row.end():]


def fix(path, rows, number, code, action):
    row = next((r for r in rows if r["no"] == number), None)
    if not row:
        sys.exit(f"접수번호 {number}이(가) 엑셀에 없습니다.")
    cells = {}
    if code:
        if parse_incident_code(code)[1]:
            sys.exit(f"장애코드 '{code}'는 유효하지 않습니다.")
        cells[f"D{row['excel_row']}"] = code.upper()
    if action:
        if parse_action_code(action)[1]:
            sys.exit(f"조치코드 '{action}'는 유효하지 않습니다.")
        cells[f"K{row['excel_row']}"] = action.upper()
    if not cells:
        sys.exit("--code 또는 --action 중 하나 이상을 지정하세요.")
    before = {"D": text(row["code"]), "K": text(row["action"])}
    backup = write_cells(path, cells)
    return {"접수번호": number, "변경": {ref: {"이전": before[ref[0]], "이후": v}
                                        for ref, v in cells.items()}, "백업": backup}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = parser.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("pending", help="원격 마지막 접수번호 이후의 등록 대상")
    p.add_argument("excel")
    p.add_argument("--remote", required=True, help="search_incidents 응답 저장 JSON")
    p.add_argument("--after", type=int, help="마지막 접수번호 직접 지정")
    p = sub.add_parser("extract", help="접수번호별 등록용 JSON 초안")
    p.add_argument("excel")
    p.add_argument("numbers", nargs="+", help="예: 1945-1950 1952 또는 1,2,3")
    p = sub.add_parser("fix", help="장애코드·조치코드 셀 수정")
    p.add_argument("excel")
    p.add_argument("number", type=int)
    p.add_argument("--code", help="D열 장애코드 (예: HL100)")
    p.add_argument("--action", help="K열 조치코드 (예: ORCH)")
    args = parser.parse_args(argv)
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(encoding="utf-8")

    rows = load_rows(args.excel)
    if args.cmd == "pending":
        result = pending(rows, args.remote, args.after)
    elif args.cmd == "extract":
        by_no = {r["no"]: r for r in rows}
        result = [build(by_no[n]) if n in by_no else {"receiptNo": n, "error": "엑셀에 없음"}
                  for n in parse_numbers(args.numbers)]
    else:
        result = fix(args.excel, rows, args.number, args.code, args.action)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
